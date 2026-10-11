"""Bounded, read-only evidence access and durable research records for adaptive loops.

No model, shell, or application-specific reasoning here. A controller chooses what to
read and which conclusions to record. Source bytes are hashed; saved excerpts remain
available even when the source changes. Records are immutable and content-addressed.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import TYPE_CHECKING

from pydantic import BaseModel

from ..contracts.evidence import (
    CITATION_SUFFIX,
    SAVED_RECORD_ID,
    CompletionDocument,
    EvidencedProposal,
    FindingDocument,
    Proposal,
    ReadRequest,
    StoredRecord,
)
from ..contracts.evidence import EvidenceError as EvidenceError  # the store's refusals keep their old import path
from ..state.response_validation import parse_json_response

if TYPE_CHECKING:  # annotations only: openpyxl stays an optional runtime dependency
    from openpyxl.workbook import Workbook


class _SourceExtract(BaseModel):
    """One source read's chunks and locator, before the record is saved."""

    chunks: list[str]
    span: dict[str, object]
    total: int


def parse_proposal(raw, *, require_finding=True):
    """Pure shared proposal format check, not registration or semantic/evidence validation."""
    return (EvidencedProposal if require_finding else Proposal).parse(raw).document


def parse_completion(raw):
    """Bounded final finding format; no source access or state mutation."""
    finding = parse_json_response(raw)
    CompletionDocument.parse(finding)
    return finding


class ResearchStore:
    def __init__(
        self, directory: Path, roots: tuple[str, ...] = (), base: str | None = None, max_source_bytes: int | None = None
    ):
        self.directory = directory
        self.roots = tuple(Path(p).resolve(strict=True) for p in roots)
        self.base = Path(base or ".").resolve(strict=True)
        if max_source_bytes is not None and (type(max_source_bytes) is not int or max_source_bytes <= 0):
            raise EvidenceError("max_source_bytes must be positive or None")
        self.max_source_bytes = max_source_bytes

    def save(self, kind: str, payload: dict) -> dict:
        if "kind" in payload or "id" in payload:
            raise EvidenceError("record kind and id are assigned by the host")
        body = {"kind": kind, **payload}
        encoded = json.dumps(body, sort_keys=True, ensure_ascii=False)
        rid = kind + "-" + hashlib.sha256(encoded.encode()).hexdigest()[:24]
        record = {"id": rid, **body}
        if self.directory.is_symlink():
            raise EvidenceError("research directory must not be a symlink")
        self.directory.mkdir(parents=True, exist_ok=True)
        path = self.directory / (rid + ".json")
        # Same record is idempotent. Never overwrite a different finding.
        if path.is_symlink():
            raise EvidenceError("research record must not be a symlink")
        if not path.exists():
            with NamedTemporaryFile(mode="w", encoding="utf-8", dir=self.directory, delete=False) as f:
                json.dump(record, f, ensure_ascii=False, indent=2)
                f.flush()
                os.fsync(f.fileno())
                tmp = Path(f.name)
            try:
                try:
                    os.link(tmp, path)  # atomic publication without overwriting an existing record
                except FileExistsError:
                    pass  # verify the concurrent/existing record below
            finally:
                tmp.unlink(missing_ok=True)
        if path.is_symlink() or json.loads(path.read_text(encoding="utf-8")) != record:
            raise EvidenceError("existing research record differs; refusing overwrite or false persistence claim")
        return record

    def recall(self, rid: str = "") -> dict:
        if not rid or rid.startswith("list:"):
            try:
                offset = int(rid[5:]) if rid else 0
                if offset < 0:
                    raise ValueError
            except ValueError as exc:
                raise EvidenceError("index cursor must be list:<nonnegative offset>") from exc
            paths = sorted(self.directory.glob("*.json"), key=lambda p: (p.stat().st_mtime_ns, p.name), reverse=True)
            records = [
                {
                    "id": p.stem,
                    "kind": (r := self.recall(p.stem))["kind"],
                    "summary": str(r.get("summary", r.get("path", r.get("hypothesis", ""))))[:200],
                }
                for p in paths[offset : offset + 20]
            ]
            return {
                "records": records,
                "total": len(paths),
                "next": f"list:{offset + 20}" if offset + 20 < len(paths) else None,
            }
        # Never permit path traversal through record ids.
        if not SAVED_RECORD_ID.fullmatch(rid):
            raise EvidenceError("invalid research record id")
        path = self.directory / (rid + ".json")
        if not path.is_file():
            raise EvidenceError(f"unknown research record {rid}")
        if path.is_symlink():
            raise EvidenceError("research record must not be a symlink")
        record = StoredRecord.parse(json.loads(path.read_text()))
        expected = record.kind + "-" + hashlib.sha256(record.canonical_body().encode()).hexdigest()[:24]
        if record.declared_id != rid or expected != rid:
            raise EvidenceError("research record content no longer matches its immutable ID")
        return record.document

    def read(self, request: dict | ReadRequest) -> dict:
        parsed, path = self._resolve_source(request)
        if path.is_dir():
            return self._list_directory(parsed, path)
        raw = self._source_bytes(path)
        suffix = path.suffix.lower()
        start = parsed.start
        end = parsed.end if parsed.end is not None else start + (2 if suffix == ".pdf" else 80) - 1
        if start < 1 or end < start:
            raise EvidenceError("start/end must select a positive one-based range")
        if suffix == ".pdf":
            extract = self._pdf_pages(raw, start, end)
        elif suffix == ".xlsx":
            if parsed.sheet is None:
                return self._xlsx_sheets(raw, path)
            extract = self._xlsx_rows(raw, parsed.sheet, start, end)
        else:
            extract = self._text_lines(raw, start, end)
        if start > extract.total:
            raise EvidenceError(f"start exceeds source length ({extract.total})")
        text = "\n".join(extract.chunks)
        return self.save(
            "source",
            {
                "path": str(path),
                "sha256": hashlib.sha256(raw).hexdigest(),
                "locator": extract.span,
                "text": text,
                "truncated": False,
            },
        )

    def _resolve_source(self, request: dict | ReadRequest) -> tuple[ReadRequest, Path]:
        """Parse the read options and resolve the path inside the roots, expanding a citation suffix."""
        parsed = ReadRequest.parse(request)
        requested = Path(parsed.path)
        # Native agents commonly cite file.py:31-33. Preserve literal filenames when
        # present; otherwise interpret this unambiguous suffix as a read locator.
        candidate = requested if requested.is_absolute() else self.base / requested
        locator = CITATION_SUFFIX.fullmatch(parsed.path)
        if not candidate.exists() and locator:
            if {"start", "end"} & parsed.written:
                raise EvidenceError("use either a citation suffix or explicit start/end, not both")
            parsed = parsed.model_copy(
                update={"path": locator[1], "start": int(locator[2]), "end": int(locator[3] or locator[2])}
            )
            requested = Path(parsed.path)
        path = (requested if requested.is_absolute() else self.base / requested).resolve(strict=True)
        if not self.roots or not any(path == root or path.is_relative_to(root) for root in self.roots):
            raise EvidenceError("source is outside --evidence-root; reads never grant themselves new roots")
        return parsed, path

    def _list_directory(self, request: ReadRequest, path: Path) -> dict:
        """Save a bounded, sorted directory listing in place of file contents."""
        # A listing has its own default end; an end written as null names no range, and is refused below.
        start, end = request.start, request.end if "end" in request.written else 80
        if end is None or start < 1 or end < start:
            raise EvidenceError("directory start/end must select a positive one-based range")
        entries = sorted(p.name + ("/" if p.is_dir() else "") for p in path.iterdir())
        return self.save(
            "source",
            {
                "path": str(path),
                "entries": entries[start - 1 : end],
                "locator": {"unit": "entry", "start": start, "end": min(end, len(entries)), "total": len(entries)},
                "text": "Directory listing, not file contents.",
            },
        )

    def _source_bytes(self, path: Path) -> bytes:
        """Read the file once, within the byte cap, so the digest identifies the exact bytes parsed."""
        if not path.is_file():
            raise EvidenceError("source must be a regular file")
        if self.max_source_bytes is not None and path.stat().st_size > self.max_source_bytes:
            raise EvidenceError("source exceeds configured max_source_bytes")
        with path.open("rb") as f:
            raw = f.read() if self.max_source_bytes is None else f.read(self.max_source_bytes + 1)
        if self.max_source_bytes is not None and len(raw) > self.max_source_bytes:
            raise EvidenceError("source exceeds configured max_source_bytes")
        return raw

    @staticmethod
    def _pdf_pages(raw: bytes, start: int, end: int) -> _SourceExtract:
        """Extract a one-based page range from PDF bytes."""
        try:
            # optional dependency: pypdf
            from pypdf import PdfReader
        except ImportError as exc:
            raise EvidenceError("PDF support needs the documents extra: pip install 'stepbook[documents]'") from exc
        reader = PdfReader(io.BytesIO(raw))
        total = len(reader.pages)
        chunks = [
            f"PAGE {n}\n{reader.pages[n - 1].extract_text() or '[No extractable text; OCR/visual inspection needed]'}"
            for n in range(start, min(end, total) + 1)
        ]
        span = {"unit": "page", "start": start, "end": min(end, total), "total": total}
        return _SourceExtract(chunks=chunks, span=span, total=total)

    @staticmethod
    def _workbook(raw: bytes) -> Workbook:
        """Open XLSX bytes read-only, or explain the missing extra."""
        try:
            # optional dependency: openpyxl
            from openpyxl import load_workbook
        except ImportError as exc:
            raise EvidenceError("XLSX support needs the documents extra: pip install 'stepbook[documents]'") from exc
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)

    def _xlsx_sheets(self, raw: bytes, path: Path) -> dict:
        """Save the sheet names of a workbook read without a sheet selected."""
        book = self._workbook(raw)
        try:
            return self.save(
                "source",
                {
                    "path": str(path),
                    "sha256": hashlib.sha256(raw).hexdigest(),
                    "sheets": book.sheetnames,
                    "text": "Choose a sheet to read rows.",
                },
            )
        finally:
            book.close()

    def _xlsx_rows(self, raw: bytes, sheet: str, start: int, end: int) -> _SourceExtract:
        """Extract a one-based row range from one sheet of a workbook."""
        book = self._workbook(raw)
        try:
            if sheet not in book.sheetnames:
                raise EvidenceError("unknown sheet; omit sheet to list available names")
            ws = book[sheet]
            total = ws.max_row or end
            chunks = [
                f"ROW {n}: " + json.dumps(list(row), default=str, ensure_ascii=False)
                for n, row in enumerate(ws.iter_rows(min_row=start, max_row=min(end, total), values_only=True), start)
            ]
            span = {"unit": "row", "sheet": sheet, "start": start, "end": min(end, total), "total": total}
            return _SourceExtract(chunks=chunks, span=span, total=total)
        finally:
            book.close()

    @staticmethod
    def _text_lines(raw: bytes, start: int, end: int) -> _SourceExtract:
        """Extract a one-based line range from UTF-8 text bytes."""
        try:
            lines = raw.decode("utf-8").splitlines()
        except UnicodeDecodeError as exc:
            raise EvidenceError("unsupported binary source; provide a text export or PDF/XLSX") from exc
        total = len(lines)
        chunks = [f"{n}: {lines[n - 1]}" for n in range(start, min(end, total) + 1)]
        span = {"unit": "line", "start": start, "end": min(end, total), "total": total}
        return _SourceExtract(chunks=chunks, span=span, total=total)

    @staticmethod
    def _validate_finding_format(finding) -> FindingDocument:
        """Check structure only. Evidence existence, authority and capture remain in record()."""
        return FindingDocument.parse(finding)

    def record(self, finding: dict) -> dict:
        document = self._validate_finding_format(finding)
        refs = [self._reference(r) for r in document.evidence]
        finding = {**finding, "evidence": refs}
        for rid in refs:
            if self.recall(rid)["kind"] not in ("source", "finding"):
                raise EvidenceError("evidence must reference a source or finding")
        if finding.get("supersedes"):
            if self.recall(finding["supersedes"])["kind"] != "finding":
                raise EvidenceError("supersedes must reference a finding")
        return self.save("finding", finding)

    def _reference(self, ref: str | ReadRequest) -> str:
        """The saved record id a finding's evidence entry stands for.

        Native file-tool users may cite a path or a locator directly. Resolve it through the exact same read
        authority and hash/snapshot contract, without an extra controller round trip just to acquire an id.
        """
        match ref:
            case str():
                return ref if SAVED_RECORD_ID.fullmatch(ref) else self.read(ReadRequest(path=ref))["id"]
            case _:
                return self.read(ref)["id"]
